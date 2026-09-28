from io import BytesIO

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def extract_text(uploaded_file):
    """
    Automatically detect the uploaded file type
    and extract its text.
    """

    file_type = uploaded_file.type
    file_data = uploaded_file.getvalue()

    # PDF

    if file_type == "application/pdf":

        reader = PdfReader(
            BytesIO(file_data)
        )

        text = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    # DOCX

    elif file_type == (
        "application/vnd.openxmlformats-officedocument"
        ".wordprocessingml.document"
    ):

        document = Document(
            BytesIO(file_data)
        )

        text = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():
                text.append(paragraph.text)

        return "\n".join(text)

    # PPTX

    elif file_type == (
        "application/vnd.openxmlformats-officedocument"
        ".presentationml.presentation"
    ):

        presentation = Presentation(
            BytesIO(file_data)
        )

        text = []

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    if shape.text.strip():
                        text.append(shape.text)

        return "\n".join(text)

    # TXT


    elif file_type == "text/plain":

        return file_data.decode(
            "utf-8",
            errors="ignore"
        )

    # Unsupported Document
    else:

        raise ValueError(
            f"Unsupported file type: {file_type}"
        )
